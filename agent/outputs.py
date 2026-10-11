"""Write the final deliverables to /output.

  catalog_agent_results.xlsx  the final output file: content, competitors, sales, data quality, run log
  content.json / competitors.json / sales_analysis.json   machine-readable copies
  report.html                 one-page human-readable report (links to /images)
  data_quality.md             every data issue found and how it was handled
  run_manifest.json           what ran, with which models, how long, what failed
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import PROJECT_ROOT

if TYPE_CHECKING:
    from .graph import PipelineState, RunContext

XLSX_NAME = "catalog_agent_results.xlsx"


def write_outputs(state: "PipelineState", ctx: "RunContext") -> dict[str, str]:
    out = ctx.settings.output_dir
    out.mkdir(parents=True, exist_ok=True)
    tables = _tables(state, ctx)

    paths = {
        "results_xlsx": _write_xlsx(out / XLSX_NAME, tables),
        "content_json": _write_json(out / "content.json", [c.model_dump() for c in state.get("content", {}).values()]),
        "competitors_json": _write_json(out / "competitors.json", {
            "summaries": [s.model_dump() for s in state.get("competitor_summaries", [])],
            "rows": [r.model_dump() for r in state.get("competitor_rows", [])],
            "our_prices": [p.model_dump() for p in state.get("our_prices", {}).values()],
        }),
        "products_json": _write_json(out / "products.json", [p.model_dump() for p in state.get("products", [])]),
        "images_json": _write_json(out / "images.json", [i.model_dump() for i in state.get("images", {}).values()]),
        "data_quality_json": _write_json(out / "data_quality.json", ctx.quality.to_records()),
        "data_quality_md": _write_text(out / "data_quality.md", _data_quality_md(state, ctx)),
    }
    if report := state.get("sales_report"):
        paths["sales_json"] = _write_json(out / "sales_analysis.json", {
            "summary": report.summary, "facts": report.facts, "treatment_notes": report.treatment_notes,
            "insights": [i.model_dump() for i in report.insights], "insights_source": report.insights_source,
            "tables": {
                "monthly": _records(report.monthly), "weekly": _records(report.weekly), "top_by_revenue": _records(report.top_by_revenue),
                "top_by_units": _records(report.top_by_units), "by_channel": _records(report.by_channel),
                "by_route": _records(report.by_route), "by_carrier": _records(report.by_carrier),
                "by_supplier": _records(report.by_supplier), "exclusions": _records(report.exclusions),
                "non_moving": _records(report.non_moving, 300), "stalled": _records(report.stalled, 300),
                "low_moving": _records(report.low_moving, 300),
            },
            "counts": {"non_moving": len(report.non_moving), "stalled": len(report.stalled),
                       "low_moving": len(report.low_moving)},
        })
    paths["report_html"] = _write_text(out / "report.html", _report_html(state, ctx))
    paths["run_manifest"] = _write_json(out / "run_manifest.json", _manifest(state, ctx, paths))
    return {k: str(Path(v).relative_to(PROJECT_ROOT)) if Path(v).is_relative_to(PROJECT_ROOT) else v
            for k, v in paths.items()}


# --- tables ---------------------------------------------------------------------------------


def _tables(state: "PipelineState", ctx: "RunContext") -> dict[str, pd.DataFrame]:
    t: dict[str, pd.DataFrame] = {"Summary": _summary_table(state, ctx)}
    t["Content"] = _content_table(state)
    t["Competitors"] = pd.DataFrame([{
        "Our SKU": r.sku, "Colour variants covered": ", ".join(r.covers_skus), "Competitor product": r.competitor_product,
        "Marketplace / website": r.marketplace, "Competitor price (USD)": r.competitor_price,
        "Our price (USD)": r.our_price, "Price difference (USD, ours - theirs)": r.price_diff_usd,
        "Price difference (%)": r.price_diff_pct, "Why comparable": r.comparability, "Rating": r.rating,
        "Reviews": r.reviews, "Link": r.link,
    } for r in state.get("competitor_rows", [])])
    t["Competitor Insights"] = pd.DataFrame([{
        "Our SKU": s.sku, "Colour variants covered": ", ".join(s.covers_skus), "Product": s.product,
        "Our price (USD)": s.our_price.price, "Our price source": s.our_price.url or s.our_price.note,
        "In stock on storefront": s.our_price.in_stock, "Competitor min": s.competitor_min,
        "Competitor median": s.competitor_median, "Competitor max": s.competitor_max,
        "Positioning": s.positioning, "Suggested launch band (unlisted SKUs)": s.suggested_price_band,
        "Competitive insight": s.insight, "Insight by": s.insight_source, "Comparables picked by": s.selection_source,
        "Search query": s.search_query, "Listings found / kept": f"{s.candidates_found} / {s.candidates_after_filter}",
        "Status": s.status + (f": {s.error}" if s.error else ""),
    } for s in state.get("competitor_summaries", [])])
    if prices := state.get("our_prices"):
        t["Our Prices"] = pd.DataFrame([{
            "SKU": p.sku, "Price (USD)": p.price, "Source": p.source, "Listing": p.url, "Listing title": p.listing_title,
            "In stock": p.in_stock, "Fetched": p.fetched_at, "Note": p.note,
        } for p in prices.values()])

    if report := state.get("sales_report"):
        kpis = pd.DataFrame([{"Metric": k.replace("_", " ").capitalize(), "Value": v} for k, v in report.summary.items()])
        notes = pd.DataFrame({"Metric": ["How sales were counted"] + [""] * (len(report.treatment_notes) - 1),
                              "Value": report.treatment_notes})
        t["Sales Summary"] = pd.concat([kpis, pd.DataFrame([{"Metric": "", "Value": ""}]), notes], ignore_index=True)
        t["Sales Insights"] = pd.DataFrame([{"#": i + 1, "Insight": x.headline, "Evidence": x.detail,
                                             "Recommendation": x.recommendation, "Written by": report.insights_source}
                                            for i, x in enumerate(report.insights)])
        t["Monthly"] = report.monthly
        t["Top SKUs by Revenue"] = report.top_by_revenue
        t["Top SKUs by Units"] = report.top_by_units
        t["Low-moving SKUs"] = report.low_moving
        t["Non-moving SKUs"] = report.non_moving
        t["Stalled SKUs"] = report.stalled
        t["Sales by Channel"] = report.by_channel
        t["Sales by Fulfilment"] = report.by_route
        t["Sales by Carrier"] = report.by_carrier
        t["Top Suppliers"] = report.by_supplier
        t["Excluded Lines"] = report.exclusions
    t["Data Quality"] = pd.DataFrame(ctx.quality.to_records())
    t["Run Log"] = pd.DataFrame(state.get("steps", []))
    return t


def _content_table(state: "PipelineState") -> pd.DataFrame:
    content = state.get("content", {})
    images = state.get("images", {})
    prices = state.get("our_prices", {})
    attr_keys: list[str] = []
    for c in content.values():
        attr_keys += [k for k in c.attributes if k not in attr_keys]
    rows = []
    for p in state.get("products", []):
        c = content.get(p.sku)
        img = images.get(p.sku)
        row = {"SKU": p.sku, "Title": c.title if c else "", "Description": c.description if c else ""}
        feats = c.key_features if c else []
        for i in range(5):
            row[f"Key Feature {i + 1}"] = feats[i] if i < len(feats) else ""
        for k in attr_keys:
            row[k] = c.attributes.get(k, "") if c else ""
        row["Lifestyle Image"] = img.path if img and img.status == "ok" else (img.note if img else "not run")
        row["Image Provider"] = img.provider if img else ""
        row["Source Image"] = p.image_url
        row["Our Price (benzara.com)"] = prices[p.sku].price if p.sku in prices else None
        row["Content QA"] = c.qa_status if c else "not run"
        row["QA Notes"] = "; ".join(c.qa_issues) if c else ""
        row["Written by"] = c.model if c else ""
        rows.append(row)
    return pd.DataFrame(rows)


def _summary_table(state: "PipelineState", ctx: "RunContext") -> pd.DataFrame:
    content = state.get("content", {})
    images = state.get("images", {})
    report = state.get("sales_report")
    rows = [
        ("Input file", ctx.settings.input_path.name),
        ("Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("SKUs processed", len(state.get("products", []))),
        ("Content: QA passed", sum(c.qa_status == "passed" for c in content.values())),
        ("Content: passed with warnings", sum(c.qa_status == "passed_with_warnings" for c in content.values())),
        ("Content: failed / error", sum(c.qa_status in ("failed", "error") for c in content.values())),
        ("Lifestyle images written", sum(i.status == "ok" for i in images.values())),
        ("Image providers used", ", ".join(sorted({i.provider for i in images.values() if i.provider})) or "-"),
        ("Models benchmarked vs competitors", len(state.get("competitor_summaries", []))),
        ("Competitor rows", len(state.get("competitor_rows", []))),
        ("Net sales (USD)", report.summary["net_sales"] if report else "-"),
        ("Units sold", report.summary["units_sold"] if report else "-"),
        ("Data issues logged", len(ctx.quality.issues)),
        ("", ""),
        ("Sheet guide", "Content = listings; Competitors / Competitor Insights = benchmarking; Sales * = analysis; "
                        "Data Quality = every anomaly and the action taken; Run Log = step status and timing"),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


# --- writers --------------------------------------------------------------------------------

WRAP_COLUMNS = {"Description", "Title", "Competitive insight", "Evidence", "Recommendation", "Insight", "Value",
                "issue", "action", "examples", "Why comparable", "Competitor product", "QA Notes", "detail", "why"}


def _write_xlsx(path: Path, tables: dict[str, pd.DataFrame]) -> str:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in tables.items():
            df = df.copy()
            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    df[col] = df[col].dt.date
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "B2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F3E46")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(df.columns, start=1):
                sample = df[col].head(200).map(str)
                longest = max([len(str(col))] + [len(s) for s in sample]) if len(sample) else len(str(col))
                wide = col in WRAP_COLUMNS or str(col).startswith("Key Feature")
                ws.column_dimensions[get_column_letter(i)].width = min(70 if wide else 40, max(10, longest + 2))
                if wide:
                    for cell in ws[get_column_letter(i)][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return str(path)


def _records(df: pd.DataFrame, limit: int | None = None) -> list[dict]:
    df = df.head(limit) if limit else df
    return json.loads(df.to_json(orient="records", date_format="iso"))


def _write_json(path: Path, data) -> str:
    path.write_text(json.dumps(data, indent=1, default=str))
    return str(path)


def _write_text(path: Path, text: str) -> str:
    path.write_text(text)
    return str(path)


def _data_quality_md(state: "PipelineState", ctx: "RunContext") -> str:
    lines = ["# Data quality report", "", f"Input: `{ctx.settings.input_path.name}`", "",
             "Every anomaly the pipeline detected and the action it took. Nothing was edited by hand.", ""]
    for sheet in ("vendor data", "sales data"):
        issues = [i for i in ctx.quality.issues if i.sheet == sheet]
        if not issues:
            continue
        lines += [f"## {sheet}", "", "| Severity | Field | Issue | Count | Action | Examples |", "|---|---|---|---|---|---|"]
        for i in issues:
            ex = "; ".join(i.examples).replace("|", "/")
            lines.append(f"| {i.severity} | {i.field} | {i.issue} | {i.count} | {i.action} | {ex} |")
        lines.append("")
    if report := state.get("sales_report"):
        lines += ["## How sales were counted", ""] + [f"- {n}" for n in report.treatment_notes] + [""]
    return "\n".join(lines)


def _manifest(state: "PipelineState", ctx: "RunContext", paths: dict[str, str]) -> dict:
    s = ctx.settings
    content = state.get("content", {}).values()
    return {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "input": {"file": s.input_path.name, "sha256": _sha256(s.input_path)},
        "settings": {
            "text_models": list(s.text_models), "fallback_text_model": s.fallback_text_model,
            "image_provider": s.image_provider, "image_model": s.image_model, "image_limit": s.image_limit,
            "skus": s.skus, "competitor_limit": s.competitor_limit, "competitors_per_sku": s.competitors_per_sku,
            "use_cache": s.use_cache,
        },
        "steps": state.get("steps", []),
        "models_used": {
            "content": sorted({c.model for c in content if c.model}),
            "images": sorted({i.provider for i in state.get("images", {}).values() if i.provider}),
            "sales_insights": state["sales_report"].insights_source if state.get("sales_report") else None,
        },
        "outputs": paths,
    }


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else ""


# --- HTML report ----------------------------------------------------------------------------


def _report_html(state: "PipelineState", ctx: "RunContext") -> str:
    from .report_html import render

    return render(state, ctx)
